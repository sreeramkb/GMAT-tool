"""
engine/error_log.py — Append-only Excel log of incorrectly-answered questions
across all sections, for post-session review/study.

Shared across sections (one workbook, "section" column distinguishes them) so
a single file accumulates every miss over time.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

ERROR_LOG_COLUMNS: List[str] = [
    "timestamp", "session_id", "section", "category", "subtopic",
    "difficulty_band", "difficulty_index", "url", "elapsed_seconds",
]

BAND_LABELS: Dict[int, str] = {
    1: "Very Easy", 2: "Easy", 3: "Easy-Medium", 4: "Medium",
    5: "Hard", 6: "Very Hard", 7: "Extreme", 8: "Abnormally Hard",
}


def _load_or_create_sheet(xlsx_path: Path):
    if xlsx_path.exists():
        wb = load_workbook(xlsx_path)
        return wb, wb.active

    wb = Workbook()
    ws = wb.active
    ws.title = "Wrong Questions"
    ws.append(ERROR_LOG_COLUMNS)
    for col_idx, header in enumerate(ERROR_LOG_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 2)
    return wb, ws


def append_wrong_questions(
    xlsx_path: Path,
    session_id: str,
    section_name: str,
    diagnostics_log: List[Dict[str, Any]],
    timestamp: Optional[str] = None,
) -> int:
    """Append one row per incorrectly-answered question in diagnostics_log. Returns rows appended."""
    xlsx_path = Path(xlsx_path)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    ts = timestamp or datetime.now(timezone.utc).isoformat()

    wb, ws = _load_or_create_sheet(xlsx_path)

    appended = 0
    for entry in diagnostics_log:
        if entry.get("was_correct"):
            continue
        band_index = entry.get("difficulty_index")
        ws.append([
            ts,
            session_id,
            section_name,
            entry.get("category"),
            entry.get("subtopic") or "",
            BAND_LABELS.get(band_index, ""),
            band_index,
            entry.get("url"),
            entry.get("elapsed_seconds"),
        ])
        appended += 1

    if appended:
        wb.save(xlsx_path)
    return appended
